"""
Generates the official Word Document (.docx) required for the Darukaa.Earth
Full-Stack Developer Hackathon submission adhering strictly to the official
Documents Submission Guidelines.
"""
from __future__ import annotations

from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, hex_color: str):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def generate_submission_docx():
    doc = Document()

    # Page setup - Standard 1 inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(36, 41, 47)

    # Document Header / Brand
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_brand = title_p.add_run("DARUKAA.EARTH")
    run_brand.font.size = Pt(13)
    run_brand.font.bold = True
    run_brand.font.color.rgb = RGBColor(16, 185, 129)  # Emerald green

    run_sub = title_p.add_run("  |  Full-Stack Developer Hackathon Submission Document")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Main Title
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(14)
    run_h1 = h1.add_run("Geospatial Intelligence Platform for Carbon & Biodiversity")
    run_h1.font.size = Pt(20)
    run_h1.font.bold = True
    run_h1.font.color.rgb = RGBColor(15, 23, 42)

    # Callout Box / Summary Card
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    c = tbl.cell(0, 0)
    set_cell_background(c, "F1F5F9")
    cp = c.paragraphs[0]
    cp.paragraph_format.space_before = Pt(6)
    cp.paragraph_format.space_after = Pt(6)
    c_run = cp.add_run(
        "Candidate Submission Document submitted via the Applied Job page as per official hackathon guidelines.\n"
        "Platform: Full-Stack Geospatial Environmental Intelligence (React 18 + Mapbox GL JS + Chart.js + "
        "FastAPI + PostgreSQL 16/PostGIS 3.4 + JWT Authentication + GitHub Actions CI/CD + Husky & lint-staged)."
    )
    c_run.font.size = Pt(9.5)
    c_run.font.italic = True
    c_run.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # =============================================================
    # 1. GitHub Repository Link
    # =============================================================
    sec1 = doc.add_paragraph()
    r = sec1.add_run("1. GitHub Repository Link")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    sec1.paragraph_format.space_before = Pt(12)
    sec1.paragraph_format.space_after = Pt(4)

    p_repo = doc.add_paragraph()
    p_repo.paragraph_format.space_after = Pt(4)
    r_link_lbl = p_repo.add_run("Repository URL: ")
    r_link_lbl.font.bold = True
    r_link = p_repo.add_run("https://github.com/thamizhan-05/darukaa-earth")
    r_link.font.bold = True
    r_link.font.color.rgb = RGBColor(2, 132, 199)

    p_desc = doc.add_paragraph()
    p_desc.paragraph_format.space_after = Pt(6)
    p_desc.add_run(
        "The repository contains the complete full-stack application code with a clean and logical commit history across "
        "infrastructure, backend models/endpoints, frontend SPA, CI/CD workflows, and architecture documentation."
    )

    # Repository Access Instructions (Explicitly required by PDF guidelines)
    p_access = doc.add_paragraph()
    p_access.paragraph_format.space_before = Pt(4)
    p_access.paragraph_format.space_after = Pt(4)
    r_acc = p_access.add_run("Repository Access Instructions (Hiring Team Access):")
    r_acc.font.bold = True

    p_acc_note = doc.add_paragraph()
    p_acc_note.paragraph_format.space_after = Pt(4)
    p_acc_note.add_run(
        "Per the challenge submission instructions, collaborator access has been granted to the hiring team accounts:"
    )

    reviewers = [
        "ankita.dasgupta@darukaa.com",
        "harsh.kumar@darukaa.com",
        "utkarsh.gauniyal@darukaa.com",
        "guneet.mutreja@darukaa.com",
    ]
    for rev in reviewers:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{rev} (Reviewer / Collaborator Access Granted)")
        r.font.color.rgb = RGBColor(71, 85, 105)

    # =============================================================
    # 2. Live Demo URL
    # =============================================================
    sec2 = doc.add_paragraph()
    r = sec2.add_run("2. Live Demo URL")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    sec2.paragraph_format.space_before = Pt(14)
    sec2.paragraph_format.space_after = Pt(4)

    p_live = doc.add_paragraph()
    p_live.paragraph_format.space_after = Pt(4)
    r_live_lbl = p_live.add_run("Publicly Accessible URL: ")
    r_live_lbl.font.bold = True
    r_live = p_live.add_run("https://daarukaearth.vercel.app")
    r_live.font.bold = True
    r_live.font.color.rgb = RGBColor(16, 185, 129)

    p_live_desc = doc.add_paragraph()
    p_live_desc.paragraph_format.space_after = Pt(6)
    p_live_desc.add_run(
        "The production web application is live and hosted on Vercel with automated continuous deployment. "
        "It features high-resolution planetary satellite basemaps (Esri World Imagery + hybrid boundary labels), "
        "interactive Mapbox GL polygon boundary drawing, real-time geodetic area calculations, Chart.js time-series analytics, "
        "and multi-tenant project management."
    )

    # =============================================================
    # 3. Brief README.md Overview
    # =============================================================
    sec3 = doc.add_paragraph()
    r = sec3.add_run("3. README.md Overview (Architecture, Schema, Local Setup, CI/CD)")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    sec3.paragraph_format.space_before = Pt(14)
    sec3.paragraph_format.space_after = Pt(6)

    # 3.1 Architecture
    p_sub1 = doc.add_paragraph()
    p_sub1.paragraph_format.space_after = Pt(2)
    r_s1 = p_sub1.add_run("A. High-Level Architecture:")
    r_s1.font.bold = True
    r_s1.font.color.rgb = RGBColor(15, 23, 42)

    arch_text = (
        "Darukaa.Earth is engineered as a clean Modular Monolith structured into clear layers:\n"
        "• Frontend: React 18 SPA built with Vite, TypeScript, and Tailwind CSS. State management via TanStack Query v5. "
        "Geospatial visualization powered by Mapbox GL JS v3 with token-free Esri World Imagery fallback and Mapbox Draw. "
        "Data visualization via Chart.js / react-chartjs-2 for all 7 environmental metrics.\n"
        "• Backend: Python 3.11 + FastAPI delivering RESTful APIs, asynchronous request handling, and Pydantic v2 schemas.\n"
        "• Security & Multi-Tenancy: JWT bearer authentication (python-jose), bcrypt password hashing (12 salt rounds), "
        "and strict role-based access control (RBAC: OWNER, ADMIN, ANALYST, VIEWER) isolating tenant organizations.\n"
        "• Spatial & Persistence: PostgreSQL 16 with PostGIS 3.4 spatial extension, GeoAlchemy2, Shapely, and Alembic migrations."
    )
    p_arch = doc.add_paragraph()
    p_arch.paragraph_format.space_after = Pt(6)
    p_arch.add_run(arch_text)

    # 3.2 Database Schema
    p_sub2 = doc.add_paragraph()
    p_sub2.paragraph_format.space_after = Pt(2)
    r_s2 = p_sub2.add_run("B. Database Schema Breakdown:")
    r_s2.font.bold = True
    r_s2.font.color.rgb = RGBColor(15, 23, 42)

    schema_tables = [
        ("users", "User credentials, hashed passwords, full names, active statuses, and timestamps."),
        ("organizations", "Multi-tenant tenant root with unique slug and metadata."),
        ("organization_members", "RBAC junction table with Role enum (OWNER, ADMIN, ANALYST, VIEWER)."),
        ("projects", "Conservation/carbon projects scoped to organization. Tracks project_type and status."),
        ("sites", "Geospatial parcels with PostGIS MULTIPOLYGON (SRID 4326) geometry, area_hectares, and GIST spatial index."),
        ("observations", "Temporal sampling points linked to site with observed timestamp and data source provenance."),
        ("observation_metrics", "7 core environmental metrics: CARBON_STOCK, CARBON_SEQUESTRATION, BIODIVERSITY_INDEX, NDVI, TREE_DENSITY, SPECIES_COUNT, CANOPY_COVER."),
        ("audit_logs", "Immutable audit trail capturing action, entity, user, organization, and JSON metadata."),
    ]
    for tbl_name, desc in schema_tables:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(2)
        r_t = p.add_run(f"{tbl_name}: ")
        r_t.font.bold = True
        r_d = p.add_run(desc)
        r_d.font.size = Pt(9.5)

    # 3.3 Local Setup
    p_sub3 = doc.add_paragraph()
    p_sub3.paragraph_format.space_before = Pt(4)
    p_sub3.paragraph_format.space_after = Pt(2)
    r_s3 = p_sub3.add_run("C. Environment Setup & Local Run Instructions:")
    r_s3.font.bold = True
    r_s3.font.color.rgb = RGBColor(15, 23, 42)

    setup_text = (
        "• Docker Compose (Recommended):\n"
        "  1. docker-compose up --build\n"
        "  2. Spawns PostGIS (port 5432), FastAPI (port 8000), and Vite frontend (port 5173).\n"
        "• Manual Setup:\n"
        "  - Backend: cd backend && pip install -e \".[dev]\" && alembic upgrade head && uvicorn app.main:app --port 8000\n"
        "  - Frontend: cd frontend && npm install && npm run dev (access at http://localhost:5173)"
    )
    p_setup = doc.add_paragraph()
    p_setup.paragraph_format.space_after = Pt(6)
    p_setup.add_run(setup_text)

    # 3.4 CI/CD Details
    p_sub4 = doc.add_paragraph()
    p_sub4.paragraph_format.space_after = Pt(2)
    r_s4 = p_sub4.add_run("D. CI/CD Pipeline & Pre-Commit Code Quality:")
    r_s4.font.bold = True
    r_s4.font.color.rgb = RGBColor(15, 23, 42)

    cicd_text = (
        "• Pre-Commit Hooks (Crucial Requirement): Implemented using Husky (.husky/pre-commit) and lint-staged. "
        "Before every commit, lint-staged automatically formats staged TypeScript/React files with Prettier and runs ESLint "
        "with zero warnings tolerance to enforce pristine code quality before code can enter Git.\n"
        "• Backend GitHub Actions (.github/workflows/backend-ci.yml): Spins up a postgis/postgis:16-3.4-alpine container, "
        "verifies Python linting/formatting with Ruff, and executes all 39 pytest tests with coverage reporting.\n"
        "• Frontend GitHub Actions (.github/workflows/frontend-ci.yml): Performs clean npm ci install, ESLint validation, "
        "TypeScript static typecheck (tsc --noEmit), production Vite build, and Vitest test suite execution."
    )
    p_cicd = doc.add_paragraph()
    p_cicd.paragraph_format.space_after = Pt(8)
    p_cicd.add_run(cicd_text)

    # =============================================================
    # 4. Reviewer Credentials & Verification Notes
    # =============================================================
    sec4 = doc.add_paragraph()
    r = sec4.add_run("4. Reviewer Credentials, Links & Operational Notes")
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 23, 42)
    sec4.paragraph_format.space_before = Pt(12)
    sec4.paragraph_format.space_after = Pt(6)

    p_cred_intro = doc.add_paragraph()
    p_cred_intro.paragraph_format.space_after = Pt(6)
    p_cred_intro.add_run(
        "A pre-seeded master administrator account is configured for instant login and testing. "
        "Reviewers can also self-register a new organization and user account on the /register screen."
    )

    cred_table = doc.add_table(rows=3, cols=3)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_table.autofit = False

    col_widths = [Inches(1.8), Inches(2.5), Inches(2.2)]
    for row in cred_table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    headers = ["Role", "Email", "Password"]
    for i, h in enumerate(headers):
        cell = cred_table.cell(0, i)
        set_cell_background(cell, "0F172A")
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)

    rows_data = [
        ("Primary Administrator", "admin@darukaa.earth", "admin1234"),
        ("New Organization Signup", "Self-register via /register", "Choose any password"),
    ]

    for r_idx, (role, email, pwd) in enumerate(rows_data, start=1):
        row = cred_table.rows[r_idx]
        for c_idx, val in enumerate([role, email, pwd]):
            cell = row.cells[c_idx]
            if r_idx % 2 == 1:
                set_cell_background(cell, "F8FAFC")
            p = cell.paragraphs[0]
            run = p.add_run(val)
            run.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Additional operational notes
    notes = [
        ("Interactive Web Database Explorer", "Accessible in the app under Settings → Database Explorer & SQL Console (/app/settings). Reviewers can inspect tables, row counts, and execute custom SELECT queries directly in the UI."),
        ("Terminal Database CLI", "Run 'py -3.11 backend/scripts/db_cli.py --interactive' for terminal SQL shell, or 'db_cli.py --tables' to view table counts."),
        ("Pre-Seeded Sample Data", "Contains real coordinates for Sundarbans Mangrove Reserve, Western Ghats Evergreen Forest, and Central Deccan Agroforestry with historical time-series observations for all 7 environmental metrics."),
        ("Interactive Polygon Drawing", "On the Map Explorer (/app/map) or within any Project site list, click 'Add Site' and draw polygon boundaries to see real-time server-side PostGIS geodetic area calculations."),
    ]
    for n_title, n_desc in notes:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r_nt = p.add_run(f"{n_title}: ")
        r_nt.font.bold = True
        r_nd = p.add_run(n_desc)
        r_nd.font.size = Pt(9.5)

    # Output path
    output_path = Path(__file__).parent.parent.parent / "DARUKAA_EARTH_SUBMISSION.docx"
    doc.save(str(output_path))
    print(f"Generated submission document successfully: {output_path}")
    return output_path


if __name__ == "__main__":
    generate_submission_docx()
